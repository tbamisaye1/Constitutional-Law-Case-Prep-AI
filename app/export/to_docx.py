"""Export a plain note (title + body) to a .docx file in memory."""

from io import BytesIO

from docx import Document


def note_to_docx(title: str, body: str) -> bytes:
    doc = Document()
    doc.add_heading(title, level=1)
    for para in body.split("\n"):
        doc.add_paragraph(para)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
